"""Portable Excel presentation of the pipeline's CSV results (no analysis rerun)."""
from __future__ import annotations

import argparse
import csv
import json
import math
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from project_config import PROJECT_ROOT

STYLE_INPUTS = [
    ("trading_style_summary.csv", "指标明细"),
    ("yearly_performance.csv", "年度表现"),
    ("monthly_performance.csv", "月度表现"),
    ("product_performance.csv", "品种分析"),
    ("product_risk_analysis.csv", "品种波动回撤"),
    ("direction_performance.csv", "方向分析"),
    ("holding_period_analysis.csv", "持仓分析"),
    ("weekday_performance.csv", "星期分析"),
    ("time_performance.csv", "时段分析"),
    ("equity_drawdown.csv", "权益回撤"),
    ("behavior_analysis.csv", "行为分析"),
]
ARBITRAGE_INPUTS = [
    ("arbitrage_candidates.csv", "套利组合清单"),
    ("arbitrage_priority_review.csv", "优先核对"),
    ("arbitrage_pair_summary.csv", "品种配对汇总"),
    ("arbitrage_leg_episodes.csv", "逐腿区间"),
    ("arbitrage_checks.csv", "检查"),
]
DARK = "17365D"
PALE = "E8EFF7"
MONEY = '#,##0.00;[Red](#,##0.00);"—"'
PERCENT = '0.0%;[Red](0.0%);"—"'
NUMBER = re.compile(r"^[-+]?(?:\d+\.?\d*|\.\d+)(?:[eE][-+]?\d+)?$")
TEXT_COLUMNS = {
    "年份", "月份", "日期", "指标", "单位", "质量", "说明", "账户", "品种", "合约",
    "方向", "候选编号", "区间ID", "A区间ID", "B区间ID", "A品种", "B品种", "A合约", "B合约",
    "平仓星期", "平仓时段", "持仓区间", "account_id", "trade_id", "contract", "product",
}


def typed_value(value: str, header: str):
    """Keep identifiers exact; convert finite numeric measures for Excel charts."""
    clean = value.strip()
    if not clean:
        return None
    if header in TEXT_COLUMNS or header.lower().endswith("_id"):
        return value
    if NUMBER.fullmatch(clean):
        # Excel supports only 15 significant decimal digits. Keep longer values as text.
        digits = re.sub(r"\D", "", clean.split("e")[0].split("E")[0]).lstrip("0")
        if len(digits) > 15:
            return value
        numeric = float(clean)
        if math.isfinite(numeric):
            return int(numeric) if numeric.is_integer() else numeric
    return value


def literal(cell, value):
    """CSV text is data, including strings starting with '='."""
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"


def read_table(path: Path):
    with path.open(encoding="utf-8-sig", newline="") as handle:
        rows = [row for row in csv.reader(handle) if any(v.strip() for v in row)]
    # The candidate producer emits a blank file for some empty result sets.
    if not rows:
        return [], []
    headers = rows[0]
    if any(not h.strip() for h in headers) or len(set(headers)) != len(headers):
        raise ValueError(f"{path.name}: blank or duplicate column names")
    for index, row in enumerate(rows[1:], 2):
        if len(row) != len(headers):
            raise ValueError(f"{path.name}: row {index} has {len(row)} values; expected {len(headers)}")
    return headers, [dict(zip(headers, row)) for row in rows[1:]]


def number_format(header):
    if header in {"胜率", "回撤比例", "较短腿重叠率", "名义金额平衡度", "平均名义金额平衡度"}:
        return PERCENT
    if any(word in header for word in ("次数", "天数", "数量", "记录数", "委托数", "明细行", "置信度")):
        return "#,##0"
    return MONEY


def format_sheet(sheet):
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "B2"
    sheet.row_dimensions[1].height = 36
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=DARK)
        cell.font = Font(name="Calibri", color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for column in sheet.columns:
        header = str(column[0].value or "")
        width = max(14, min(48, max(len(str(c.value or "")) * 1.6 for c in column[:101]) + 2))
        sheet.column_dimensions[column[0].column_letter].width = width
        for cell in column[1:]:
            cell.font = Font(name="Calibri", size=11)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, (int, float)):
                cell.number_format = number_format(header)
        if sheet.max_row > 1:
            area = f"{column[0].column_letter}2:{column[0].column_letter}{sheet.max_row}"
            sheet.conditional_formatting.add(area, CellIsRule(
                operator="lessThan", formula=["0"], font=Font(color="C00000")))
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"


def import_tables(workbook, data_dir, inputs):
    missing = [filename for filename, _ in inputs if not (data_dir / filename).is_file()]
    if missing:
        raise ValueError("Missing input files: " + ", ".join(missing) + ". Run the analysis stage first.")
    tables = {}
    for index, (filename, name) in enumerate(inputs, 1):
        headers, rows = read_table(data_dir / filename)
        sheet = workbook.create_sheet(name)
        for col, header in enumerate(headers, 1):
            literal(sheet.cell(1, col), header)
        for row_index, row in enumerate(rows, 2):
            for col, header in enumerate(headers, 1):
                literal(sheet.cell(row_index, col), typed_value(row[header], header))
        if not headers:
            sheet["A1"] = "无数据"
        elif rows:
            table = Table(displayName=f"ReportData{index}", ref=sheet.dimensions)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            sheet.add_table(table)
        else:
            sheet.cell(2, 1, "无数据 / No data")
        format_sheet(sheet)
        if "单位" in headers and "数值" in headers:
            for row_index, row in enumerate(rows, 2):
                sheet.cell(row_index, headers.index("数值") + 1).number_format = {
                    "%": PERCENT, "元": MONEY, "笔": "#,##0", "次": "#,##0"
                }.get(row["单位"], "0.00")
        tables[name] = (headers, rows)
    return tables


def require_columns(tables, name, columns):
    headers, rows = tables[name]
    missing = set(columns) - set(headers)
    if missing:
        raise ValueError(f"{name}: missing columns: {', '.join(sorted(missing))}")
    return headers, rows


def reference(sheet, headers, field, row):
    col = get_column_letter(headers.index(field) + 1)
    return f"'{sheet}'!{col}{row}"


def linked_value(cell, ref):
    # Preserve an absent statistic instead of displaying a plausible zero.
    cell.value = f'=IF({ref}="","",{ref})'


def cover_sheet(workbook, title, period, note):
    sheet = workbook.active
    sheet.title = "总览" if "风格" in title else "使用说明"
    sheet.merge_cells("A1:H2")
    sheet["A1"] = title
    sheet["A1"].font = Font(name="Calibri", size=22, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor=DARK)
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 26
    sheet.row_dimensions[2].height = 20
    sheet.merge_cells("A3:H3")
    sheet["A3"] = f"数据期间：{period}"
    sheet.merge_cells("A5:H7")
    sheet["A5"] = note
    sheet["A5"].alignment = Alignment(wrap_text=True, vertical="center")
    sheet["A5"].fill = PatternFill("solid", fgColor="FFF2CC")
    for row in (5, 6, 7):
        sheet.row_dimensions[row].height = 24
    for col in range(1, 9):
        sheet.column_dimensions[get_column_letter(col)].width = 18
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A9"
    return sheet


def chart(workbook, tables, source, category, fields, title, anchor, *, target=None, line=False, limit=None):
    headers, rows = require_columns(tables, source, [category, *fields])
    count = min(len(rows), limit) if limit else len(rows)
    if not count:
        return
    sheet = workbook[source]
    obj = LineChart() if line else BarChart()
    obj.title, obj.style = title, 13
    obj.width, obj.height = 23, 11
    obj.y_axis.numFmt = "#,##0"
    for field in fields:
        col = headers.index(field) + 1
        obj.add_data(Reference(sheet, min_col=col, max_col=col, min_row=1, max_row=count + 1), titles_from_data=True)
    col = headers.index(category) + 1
    obj.set_categories(Reference(sheet, min_col=col, max_col=col, min_row=2, max_row=count + 1))
    if len(fields) == 1:
        obj.legend = None
    workbook[target or source].add_chart(obj, anchor)


def build_style(data_dir: Path):
    workbook = Workbook()
    tables = import_tables(workbook, data_dir, STYLE_INPUTS)
    _, equity = require_columns(tables, "权益回撤", ["日期"])
    dates = sorted(row["日期"] for row in equity if row["日期"])
    period = f"{dates[0]} 至 {dates[-1]}" if dates else "无数据"
    cover = cover_sheet(workbook, "交易风格分析报告", period,
        "TXT 没有精确成交时间，盘中时段和行为分析仅使用精确时间样本。"
        "非月末权益不含无法反推的浮动盈亏，回撤为估算值。品种回撤按已实现净收益计算。"
        "账户总额与逐品种归因可能因在途开仓费等存在差异。")
    headers, metrics = require_columns(tables, "指标明细", ["指标", "数值", "单位"])
    cover.append([])
    for col, value in enumerate(["指标", "数值", "单位"], 1):
        cover.cell(9, col, value)
    # Locate metric rows by name instead of assuming a fixed input row ordering.
    lookup = {row["指标"]: (i, row) for i, row in enumerate(metrics, 2)}
    for output_row, metric in enumerate([
        "累计账户实际收益", "累计已实现净收益", "期末权益", "胜率", "利润因子", "前10笔盈利集中度"
    ], 10):
        cover.cell(output_row, 1, metric)
        if metric not in lookup:
            raise ValueError(f"指标明细: missing metric {metric}")
        source_row, data = lookup[metric]
        linked_value(cover.cell(output_row, 2), reference("指标明细", headers, "数值", source_row))
        literal(cover.cell(output_row, 3), data["单位"])
        cover.cell(output_row, 2).number_format = PERCENT if data["单位"] == "%" else MONEY
        cover.row_dimensions[output_row].height = 26
    cover.column_dimensions["A"].width = 28
    cover.column_dimensions["B"].width = 23
    chart(workbook, tables, "年度表现", "年份", ["已实现净收益", "账户实际收益"], "年度收益对比（元）", "A18", target=cover.title)
    chart(workbook, tables, "品种分析", "品种", ["净收益"], "品种净收益（输入排序前10项，元）", "A40", target=cover.title, limit=10)
    chart(workbook, tables, "月度表现", "月份", ["期末权益"], "月末权益（元）", "P2", line=True)
    chart(workbook, tables, "品种波动回撤", "品种", ["最大回撤金额"], "品种回撤（输入排序前15项，元）", "O2", limit=15)
    add_sources(cover, STYLE_INPUTS, 63)
    return workbook, tables


def add_sources(sheet, inputs, start):
    sheet.cell(start, 1, "工作表")
    sheet.cell(start, 3, "输入文件")
    for row, (filename, name) in enumerate(inputs, start + 1):
        sheet.cell(row, 1, name)
        sheet.merge_cells(start_row=row, start_column=3, end_row=row, end_column=8)
        sheet.cell(row, 3, filename)
    for cell in (sheet.cell(start, 1), sheet.cell(start, 3)):
        cell.font = Font(bold=True, color=DARK)


def build_arbitrage(data_dir: Path):
    workbook = Workbook()
    tables = import_tables(workbook, data_dir, ARBITRAGE_INPUTS)
    metadata_path = data_dir / "arbitrage_metadata.json"
    if not metadata_path.is_file():
        raise ValueError("Missing arbitrage_metadata.json. Run identify_arbitrage.py first.")
    metadata = json.loads(metadata_path.read_text(encoding="utf-8-sig"))
    required = ["date_start", "date_end", "episodes", "candidates", "priority_pairs", "high", "medium", "low", "minimum_score"]
    if any(key not in metadata for key in required):
        raise ValueError("arbitrage_metadata.json: incomplete metadata")
    cover = cover_sheet(workbook, "对冲套利候选识别报告",
        f"{metadata['date_start']} 至 {metadata['date_end']}",
        "账单没有套利策略编号，本报告展示候选组合，不能确认交易意图。"
        "同一持仓腿可能进入多个候选，全部候选的组合净收益不能直接相加。"
        "优先核对页由上游贪心配对选取，可能遗漏多腿组合。TXT 的时间精度受限。")
    for row, (label, key) in enumerate([
        ("持仓区间", "episodes"), ("全部候选", "candidates"), ("优先核对配对", "priority_pairs"),
        ("高置信度", "high"), ("中置信度", "medium"), ("低置信度", "low"), ("最低识别得分", "minimum_score")
    ], 10):
        cover.cell(row, 1, label)
        literal(cover.cell(row, 3), metadata[key])
        cover.row_dimensions[row].height = 25
    cover.merge_cells("A19:H21")
    cover["A19"] = "筛选依据：不同合约、相反方向和持仓重叠；评分考虑品种关系、开平日期协调、重叠比例及名义金额平衡。具体规则以 identify_arbitrage.py 为准。"
    cover["A19"].alignment = Alignment(wrap_text=True, vertical="center")
    add_sources(cover, ARBITRAGE_INPUTS + [("arbitrage_metadata.json", "报告概况")], 24)
    for name in ("套利组合清单", "优先核对"):
        sheet = workbook[name]
        headers, rows = tables[name]
        if "置信度" in headers:
            for i, data in enumerate(rows, 2):
                color = {"高": "E2F0D9", "中": "FFF2CC", "低": "E7E6E6"}.get(data["置信度"], "FFFFFF")
                sheet.cell(i, headers.index("置信度") + 1).fill = PatternFill("solid", fgColor=color)
        sheet.freeze_panes = "F2"
    return workbook, tables


def export_previews(tables, kind, destination, *, synthetic=False):
    """Standalone chart summaries, not screenshots of Excel's rendered worksheets."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    fonts = {font.name for font in font_manager.fontManager.ttflist}
    for name in ("Microsoft YaHei", "SimHei", "Noto Sans CJK SC", "WenQuanYi Zen Hei"):
        if name in fonts:
            plt.rcParams["font.sans-serif"] = [name, "DejaVu Sans"]
            break
    plt.rcParams["axes.unicode_minus"] = False
    destination.mkdir(parents=True, exist_ok=True)
    specs = [
        ("年度表现", "年份", "账户实际收益", "annual_profit", False),
        ("月度表现", "月份", "期末权益", "monthly_equity", True),
        ("品种分析", "品种", "净收益", "product_profit", False),
        ("品种波动回撤", "品种", "最大回撤金额", "product_drawdown", False),
    ] if kind == "style" else [("品种配对汇总", "品种组合", "候选数量", "candidate_counts", False)]
    paths = []
    for source, category, field, filename, line in specs:
        headers, rows = tables[source]
        fig, ax = plt.subplots(figsize=(12, 6), layout="constrained")
        if synthetic:
            fig.suptitle("SYNTHETIC DEMO - NOT HISTORICAL PERFORMANCE", color="#9C5700")
        if rows:
            require_columns(tables, source, [category, field])
            selected = rows if line or source == "年度表现" else rows[:15]
            labels = [str(r[category]) for r in selected]
            values = [float(r[field]) if r[field].strip() else float("nan") for r in selected]
            if line:
                ax.plot(range(len(labels)), values, color="#0F6B78")
            else:
                bars = ax.bar(range(len(labels)), values, color="#17365D")
                ax.bar_label(bars, labels=[f"{v:,.2f}" if math.isfinite(v) else "" for v in values], padding=3)
                ax.margins(y=.15)
                if all(v == 0 for v in values):
                    ax.set_ylim(0, 1)
                if field == "候选数量":
                    from matplotlib.ticker import MaxNLocator
                    ax.yaxis.set_major_locator(MaxNLocator(integer=True))
            stride = max(1, math.ceil(len(labels) / 20))
            ax.set_xticks(list(range(0, len(labels), stride)), labels[::stride], rotation=45, ha="right")
            ax.set_title(f"{source} / {field}")
            ax.grid(axis="y", alpha=0.2)
        else:
            ax.text(.5, .5, "No data", ha="center", va="center", transform=ax.transAxes)
            ax.set_axis_off()
        path = destination / f"{filename}.png"
        fig.savefig(path, dpi=140)
        plt.close(fig)
        paths.append(str(path))
    return paths


def report_cli(kind):
    parser = argparse.ArgumentParser(description=f"Build the {kind} Excel report from analysis CSV files.")
    parser.add_argument("--data-dir", type=Path, default=PROJECT_ROOT / "output")
    parser.add_argument("--output", type=Path, help="Output .xlsx path (default: DATA_DIR/reports/<kind>_report.xlsx)")
    parser.add_argument("--preview-dir", type=Path, help="Optionally export PNG chart summaries")
    parser.add_argument("--synthetic", action="store_true", help="Label the workbook and PNGs as synthetic demo data")
    args = parser.parse_args()
    output = args.output or args.data_dir / "reports" / f"{kind}_report.xlsx"
    if output.suffix.lower() != ".xlsx":
        parser.error("--output must end in .xlsx")
    try:
        workbook, tables = (build_style if kind == "style" else build_arbitrage)(args.data_dir)
        if args.synthetic:
            workbook.active["A1"] = "合成演示 / " + workbook.active["A1"].value
        output.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output)
        previews = export_previews(tables, kind, args.preview_dir, synthetic=args.synthetic) if args.preview_dir else []
    except (ValueError, OSError) as exc:
        parser.exit(1, f"Report generation failed: {exc}\n")
    print(json.dumps({"output": str(output.resolve()), "sheets": workbook.sheetnames, "previews": previews}, ensure_ascii=False))
