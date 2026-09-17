"""Portable report integration tests using only generated synthetic inputs."""
import csv
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from demo_reports import prepare_analysis
from report_workbook import build_style, build_arbitrage, read_table, typed_value


class ReportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.fixture = tempfile.TemporaryDirectory()
        cls.data = Path(cls.fixture.name)
        prepare_analysis(cls.data)

    @classmethod
    def tearDownClass(cls):
        cls.fixture.cleanup()

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.local = Path(self.temp.name)
        shutil.copytree(self.data, self.local, dirs_exist_ok=True)

    def test_style_sources_dates_and_chart_extents(self):
        workbook, _ = build_style(self.local)
        output = self.local / "style.xlsx"
        workbook.save(output)
        with output.open("rb") as handle:
            saved = load_workbook(handle)
            self.assertEqual(len(saved.sheetnames), 12)
            self.assertEqual(saved["总览"]["A3"].value, "数据期间：2024-01-01 至 2024-02-29")
            self.assertEqual(saved["总览"]["B10"].value, '=IF(\'指标明细\'!B3="","",\'指标明细\'!B3)')
            self.assertEqual(saved["指标明细"]["B3"].value, 384)
            self.assertEqual(saved["月度表现"]._charts[0].series[0].val.numRef.f, "'月度表现'!$H$2:$H$3")
            self.assertEqual(saved["品种分析"]["F2"].number_format, '0.0%;[Red](0.0%);"—"')
            self.assertEqual(saved["时段分析"].max_row, 3)
            saved.close()

    def test_reordered_columns_and_metrics_keep_correct_references(self):
        path = self.local / "trading_style_summary.csv"
        headers, rows = read_table(path)
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=list(reversed(headers)))
            writer.writeheader()
            writer.writerows(reversed(rows))
        workbook, _ = build_style(self.local)
        row_index = next(i for i, row in enumerate(reversed(rows), 2) if row["指标"] == "累计账户实际收益")
        self.assertIn(f"'指标明细'!D{row_index}", workbook["总览"]["B10"].value)

    def test_arbitrage_preserves_candidates_and_identifiers(self):
        path = self.local / "arbitrage_leg_episodes.csv"
        headers, rows = read_table(path)
        rows[0]["账户"] = "001234567890123456789"
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=headers)
            writer.writeheader()
            writer.writerows(rows)
        workbook, tables = build_arbitrage(self.local)
        self.assertEqual(len(workbook.sheetnames), 6)
        self.assertEqual(workbook["逐腿区间"]["B2"].value, "001234567890123456789")
        self.assertEqual(len(tables["套利组合清单"][1]), 2)
        self.assertEqual(workbook["使用说明"]["C11"].value, 2)

    def test_empty_candidates_and_header_only_time_analysis(self):
        for filename in ("arbitrage_pair_summary.csv", "arbitrage_leg_episodes.csv"):
            (self.local / filename).write_text("\n", encoding="utf-8-sig")
        for filename in ("arbitrage_candidates.csv", "arbitrage_priority_review.csv", "time_performance.csv"):
            headers, _ = read_table(self.local / filename)
            with (self.local / filename).open("w", encoding="utf-8-sig", newline="") as handle:
                csv.writer(handle).writerow(headers)
        metadata_path = self.local / "arbitrage_metadata.json"
        metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
        for key in ("episodes", "candidates", "priority_pairs", "high", "medium", "low"):
            metadata[key] = 0
        metadata_path.write_text(json.dumps(metadata), encoding="utf-8")
        workbook, _ = build_arbitrage(self.local)
        workbook.save(self.local / "empty.xlsx")
        self.assertEqual(workbook["品种配对汇总"]["A1"].value, "无数据")
        style, _ = build_style(self.local)
        self.assertEqual(style["时段分析"]["A2"].value, "无数据 / No data")

    def test_missing_input_reports_actionable_error(self):
        (self.local / "yearly_performance.csv").unlink()
        with self.assertRaisesRegex(ValueError, "yearly_performance.csv"):
            build_style(self.local)

    def test_literal_csv_text_does_not_become_formula(self):
        path = self.local / "behavior_analysis.csv"
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(["指标", "数值", "单位", "质量", "说明"])
            writer.writerow(["example", "1", "次", "合成", '=HYPERLINK("https://example.invalid","text")'])
        workbook, _ = build_style(self.local)
        workbook.save(self.local / "literal.xlsx")
        with (self.local / "literal.xlsx").open("rb") as handle:
            saved = load_workbook(handle)
            self.assertEqual(saved["行为分析"]["E2"].data_type, "s")
            saved.close()
        self.assertEqual(typed_value("001234", "账户"), "001234")
        self.assertEqual(typed_value("1234567890123456789", "成交额"), "1234567890123456789")
        self.assertEqual(typed_value("-3.5", "净收益"), -3.5)

    def test_new_year_and_reordered_chart_columns_are_dynamic(self):
        path = self.local / "yearly_performance.csv"
        headers, rows = read_table(path)
        extra = dict(rows[0], 年份="2025")
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=list(reversed(headers)))
            writer.writeheader()
            writer.writerows([*rows, extra])
        workbook, _ = build_style(self.local)
        self.assertEqual(workbook["总览"]._charts[0].series[0].val.numRef.f, "'年度表现'!$E$2:$E$3")

    def test_cli_outputs_and_error_exit(self):
        script = Path(__file__).resolve().parents[1] / "scripts" / "build_style_report.py"
        result = subprocess.run([sys.executable, str(script), "--data-dir", str(self.local), "--synthetic"], capture_output=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertTrue((self.local / "reports" / "style_report.xlsx").is_file())
        with (self.local / "reports" / "style_report.xlsx").open("rb") as handle:
            saved = load_workbook(handle)
            self.assertTrue(saved.active["A1"].value.startswith("合成演示"))
            saved.close()
        result = subprocess.run([sys.executable, str(script), "--data-dir", str(self.local / "missing")], capture_output=True)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn(b"Run the analysis stage first", result.stderr)

    def test_supported_xlsx_and_txt_statements_to_reports(self):
        """Exercise documented raw layouts through both report entry points."""
        scripts = Path(__file__).resolve().parents[1] / "scripts"
        _, all_rows = read_table(self.local / "transactions.csv")
        executions = [row for row in all_rows if row["statement_month"] == "2024-01"]
        for extension in ("xlsx", "txt"):
            with self.subTest(extension=extension):
                raw = self.local / extension / "raw"
                output = self.local / extension / "output"
                raw.mkdir(parents=True)
                source = raw / f"000001-202401账单.{extension}"
                if extension == "xlsx":
                    book = Workbook()
                    sheet = book.active
                    sheet.title = "客户交易结算月报"
                    sheet.append(["交易月份", "2024-01"])
                    sheet.append(["客户期货期权内部资金账户", "000001"])
                    for i, (label, value) in enumerate([
                        ("上月结存", 100000), ("当月存取合计", 0), ("当月盈亏", 200),
                        ("当月手续费", 8), ("当月结存", 100192), ("客户权益", 100192), ("浮动盈亏", 0),
                    ], 9):
                        sheet.cell(i, 1, label)
                        sheet.cell(i, 2, value)
                    sheet = book.create_sheet("成交明细")
                    columns = {
                        "交易日期": "trade_date", "合约": "contract", "成交序号": "trade_id",
                        "成交时间": "trade_time", "买/卖": "side_cn", "成交价": "price",
                        "手数": "quantity", "开/平": "position_effect_cn", "手续费": "fee",
                        "平仓盈亏": "realized_pnl", "实际成交日期": "actual_trade_date", "成交额": "turnover",
                    }
                    sheet.append(list(columns))
                    for transaction in executions:
                        sheet.append([transaction[field] for field in columns.values()])
                    book.save(source)
                    book.close()
                else:
                    lines = [f"{label}: {value}" for label, value in [
                        ("期初结存", 100000), ("期间入金", 0), ("期间出金", 0),
                        ("期间出入金合计", 0), ("平仓盈亏", 200), ("期间手续费", 8),
                        ("期末结存", 100192), ("资金权益", 100192), ("浮动持仓盈亏", 0),
                    ]]
                    lines.append("成交日期 成交时间")
                    for t in executions:
                        lines.append(" ".join([
                            t["trade_date"].replace("-", ""), t["trade_time"], "DCE", t["contract"],
                            t["trade_id"], t["side_cn"], "投机", t["price"], t["quantity"], t["turnover"],
                            t["position_effect_cn"], t["fee"], t["realized_pnl"], "0", "0",
                        ]))
                    source.write_text("\n".join(lines), encoding="gb18030")
                commands = [
                    ["parse_statements.py", "--source", str(raw), "--output", str(output), "--year", "2024"],
                    ["match_trades.py", "--transactions", str(output / "transactions.csv"),
                     "--output", str(output / "completed_trades.csv"), "--errors", str(output / "parsing_errors.csv"),
                     "--source-root", str(raw)],
                    ["analyze_trading_style.py", "--input-dir", str(output), "--output-dir", str(output)],
                    ["identify_arbitrage.py", "--completed-trades", str(output / "completed_trades.csv"), "--output-dir", str(output)],
                    ["build_style_report.py", "--data-dir", str(output), "--synthetic"],
                    ["build_arbitrage_report.py", "--data-dir", str(output), "--synthetic"],
                ]
                for command, *args in commands:
                    result = subprocess.run([sys.executable, str(scripts / command), *args], capture_output=True)
                    self.assertEqual(result.returncode, 0, result.stderr)
                _, parsed = read_table(output / "transactions.csv")
                self.assertEqual(len(parsed), 4)
                self.assertTrue(all(row["account_id"] == "000001" for row in parsed))
                _, errors = read_table(output / "parsing_errors.csv")
                self.assertEqual(errors, [])
                for name, sheets in [("style_report.xlsx", 12), ("arbitrage_report.xlsx", 6)]:
                    with (output / "reports" / name).open("rb") as handle:
                        workbook = load_workbook(handle)
                        self.assertEqual(len(workbook.sheetnames), sheets)
                        if sheets == 12:
                            self.assertEqual(workbook["指标明细"]["B3"].value, 192)
                        workbook.close()


if __name__ == "__main__":
    unittest.main()
